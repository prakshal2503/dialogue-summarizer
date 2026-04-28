"""
Views for Dialogue Summarizer
"""
import os
import json
import hashlib
import datetime
import logging
from functools import wraps

from django.shortcuts import render, redirect
from django.http import JsonResponse, Http404
from django.views.decorators.http import require_http_methods, require_POST
from django.views.decorators.csrf import csrf_protect
from django.contrib import messages
from django.conf import settings

from .db import (
    init_db, authenticate_user, create_user, get_user_by_id,
    get_user_by_username, get_all_users, count_users,
    save_upload_history, get_user_history, count_user_history,
    get_all_history, get_system_stats, record_login_attempt,
    is_ip_locked, toggle_user_status, delete_user, update_user_profile,
    compute_sha256, get_db,
    delete_history_item, clear_user_history, delete_own_account,
    save_audit_log, get_all_audit_logs, count_audit_logs
)
from .forms import LoginForm, RegisterForm, UploadForm, ProfileForm, verify_recaptcha

logger = logging.getLogger('core.security')


# ─── Decorators ──────────────────────────────────────────────────────────────

def login_required(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.session.get('user_id'):
            return redirect(f"{settings.LOGIN_URL}?next={request.path}")
        return view_func(request, *args, **kwargs)
    return wrapper


def admin_required(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.session.get('user_id'):
            return redirect(settings.LOGIN_URL)
        if request.session.get('user_role') != 'admin':
            messages.error(request, "Access denied. Admin privileges required.")
            return redirect('/dashboard/')
        return view_func(request, *args, **kwargs)
    return wrapper


def get_client_ip(request):
    x_forwarded = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded:
        return x_forwarded.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR', '127.0.0.1')


# ─── Auth Views ──────────────────────────────────────────────────────────────

@csrf_protect
def login_view(request):
    if request.session.get('user_id'):
        return redirect('/dashboard/')

    form = LoginForm()
    error = None

    if request.method == 'POST':
        form = LoginForm(request.POST)
        ip = get_client_ip(request)

        if is_ip_locked(ip):
            error = "Too many failed attempts. Please wait 15 minutes."
        elif form.is_valid():
            username = form.cleaned_data['username']
            password = form.cleaned_data['password']
            recaptcha_token = request.POST.get('g-recaptcha-response', '')

            if not verify_recaptcha(recaptcha_token, ip):
                error = "Please complete the CAPTCHA verification."
            else:
                user = authenticate_user(username, password)
                if user:
                    record_login_attempt(ip, username, success=True)
                    request.session['user_id'] = str(user['_id'])
                    request.session['username'] = user['username']
                    request.session['user_role'] = user['role']
                    request.session['full_name'] = user['profile'].get('full_name', '')
                    request.session.set_expiry(settings.SESSION_COOKIE_AGE)
                    next_url = request.GET.get('next', '/dashboard/')
                    return redirect(next_url)
                else:
                    record_login_attempt(ip, username, success=False)
                    error = "Invalid credentials. Please try again."
        else:
            error = "Please fill in all fields."

    return render(request, 'core/login.html', {
        'form': form,
        'error': error,
        'recaptcha_site_key': settings.RECAPTCHA_PUBLIC_KEY,
    })


@csrf_protect
def register_view(request):
    if request.session.get('user_id'):
        return redirect('/dashboard/')

    form = RegisterForm()
    error = None

    if request.method == 'POST':
        form = RegisterForm(request.POST)
        ip = get_client_ip(request)
        recaptcha_token = request.POST.get('g-recaptcha-response', '')

        if not verify_recaptcha(recaptcha_token, ip):
            error = "Please complete the CAPTCHA verification."
        elif form.is_valid():
            result = create_user(
                username=form.cleaned_data['username'],
                email=form.cleaned_data['email'],
                password=form.cleaned_data['password'],
                full_name=form.cleaned_data['full_name'],
            )
            if result['success']:
                user = result['user']
                request.session['user_id'] = str(user['_id'])
                request.session['username'] = user['username']
                request.session['user_role'] = user['role']
                request.session['full_name'] = user['profile'].get('full_name', '')
                messages.success(request, "Account created successfully! Welcome aboard.")
                return redirect('/dashboard/')
            else:
                error = result['error']
        else:
            errors = []
            for field, errs in form.errors.items():
                for e in errs:
                    errors.append(e)
            error = " ".join(errors)

    return render(request, 'core/register.html', {
        'form': form,
        'error': error,
        'recaptcha_site_key': settings.RECAPTCHA_PUBLIC_KEY,
    })


def logout_view(request):
    request.session.flush()
    return redirect('/login/')


# ─── Dashboard Views ──────────────────────────────────────────────────────────

@login_required
def dashboard(request):
    user_id = request.session['user_id']
    role = request.session.get('user_role')

    if role == 'admin':
        return redirect('/admin-dashboard/')

    user = get_user_by_id(user_id)
    if not user:
        request.session.flush()
        return redirect('/login/')

    history = get_user_history(user_id, limit=10)
    total_history = count_user_history(user_id)

    for item in history:
        item['id'] = str(item['_id'])

    return render(request, 'core/user_dashboard.html', {
        'user': user,
        'history': history,
        'total_history': total_history,
        'user_id': user_id,
    })


@admin_required
def admin_dashboard(request):
    stats = get_system_stats()
    recent_users = get_all_users(limit=5)
    for u in recent_users:
        u['id'] = str(u['_id'])
    recent_activity = get_all_history(limit=10)
    for item in recent_activity:
        item['id'] = str(item['_id'])

    return render(request, 'core/admin_dashboard.html', {
        'stats': stats,
        'recent_users': recent_users,
        'recent_activity': recent_activity,
    })


@admin_required
def admin_users(request):
    page = int(request.GET.get('page', 1))
    limit = settings.ITEMS_PER_PAGE
    skip = (page - 1) * limit
    users = get_all_users(skip=skip, limit=limit)
    total = count_users()
    total_pages = (total + limit - 1) // limit

    for u in users:
        u['id'] = str(u['_id'])

    return render(request, 'core/admin_users.html', {
        'users': users,
        'total': total,
        'page': page,
        'total_pages': total_pages,
    })


@admin_required
def admin_activity(request):
    page = int(request.GET.get('page', 1))
    limit = 20
    skip = (page - 1) * limit
    activity = get_all_history(skip=skip, limit=limit)
    for item in activity:
        item['id'] = str(item['_id'])

    return render(request, 'core/admin_activity.html', {
        'activity': activity,
        'page': page,
    })


@admin_required
@require_POST
def admin_toggle_user(request, user_id):
    new_status = toggle_user_status(user_id)
    status_text = 'activated' if new_status else 'deactivated'
    messages.success(request, f"User {status_text} successfully.")
    return redirect('/admin/users/')


@admin_required
@require_POST
def admin_delete_user(request, user_id):
    if delete_user(user_id):
        messages.success(request, "User deleted successfully.")
    else:
        messages.error(request, "Could not delete user.")
    return redirect('/admin/users/')


# ─── Upload / Summarize Views ─────────────────────────────────────────────────

@login_required
@csrf_protect
def upload_file(request):
    user_id = request.session['user_id']
    form = UploadForm()
    error = None
    success_msg = None

    if request.method == 'POST':
        form = UploadForm(request.POST, request.FILES)
        if form.is_valid():
            uploaded_file = request.FILES['file']
            description = form.cleaned_data.get('description', '')
            content = uploaded_file.read()

            file_hash = compute_sha256(content)

            upload_dir = os.path.join(settings.MEDIA_ROOT, 'uploads', user_id)
            os.makedirs(upload_dir, exist_ok=True)
            safe_name = f"{file_hash[:8]}_{uploaded_file.name}"
            file_path = os.path.join(upload_dir, safe_name)
            with open(file_path, 'wb') as f:
                f.write(content)

            save_upload_history(
                user_id=user_id,
                filename=uploaded_file.name,
                file_hash=file_hash,
                file_size=uploaded_file.size,
                action='upload',
                result=description or f"Uploaded: {uploaded_file.name}"
            )

            success_msg = f"File '{uploaded_file.name}' uploaded successfully! SHA-256: {file_hash[:16]}..."
        else:
            error = str(list(form.errors.values())[0][0]) if form.errors else "Upload failed."

    history = get_user_history(user_id, limit=5)
    for item in history:
        item['id'] = str(item['_id'])
    return render(request, 'core/upload.html', {
        'form': form,
        'error': error,
        'success_msg': success_msg,
        'history': history,
    })


@login_required
@csrf_protect
def summarize_view(request):
    user_id = request.session['user_id']
    summary = None
    error = None
    extracted_text = ''
    input_filename = 'dialogue_text'
    input_mode = 'text'

    if request.method == 'POST':
        input_mode = request.POST.get('input_mode', 'text')

        if input_mode == 'file' and request.FILES.get('input_file'):
            uploaded_file = request.FILES['input_file']
            input_filename = uploaded_file.name

            extracted_text, extract_error = _extract_text_from_file(uploaded_file)

            if extract_error:
                error = f'Could not read "{input_filename}": {extract_error}'
            elif not extracted_text.strip():
                error = f'No text could be extracted from "{input_filename}". Try pasting the text directly.'
            else:
                summary = _generate_summary(extracted_text)
                text_hash = compute_sha256(extracted_text.encode())
                save_upload_history(
                    user_id=user_id,
                    filename=input_filename,
                    file_hash=text_hash,
                    file_size=len(extracted_text),
                    action='summarize',
                    result=summary[:200] + '...' if len(summary) > 200 else summary
                )
            extracted_text = extracted_text[:500]

        else:
            text = request.POST.get('dialogue_text', '').strip()
            if not text:
                error = 'Please provide dialogue text to summarize.'
            elif len(text) < 10:
                error = 'Text too short to summarize.'
            else:
                summary = _generate_summary(text)
                text_hash = compute_sha256(text.encode())
                save_upload_history(
                    user_id=user_id,
                    filename='dialogue_text',
                    file_hash=text_hash,
                    file_size=len(text),
                    action='summarize',
                    result=summary[:200] + '...' if len(summary) > 200 else summary
                )

    history = get_user_history(user_id, limit=5)
    for item in history:
        item['id'] = str(item['_id'])
    return render(request, 'core/summarize.html', {
        'summary': summary,
        'error': error,
        'history': history,
        'dialogue_text': request.POST.get('dialogue_text', ''),
        'extracted_text': extracted_text,
        'input_mode': input_mode,
        'input_filename': input_filename,
    })


def _extract_text_from_file(uploaded_file) -> tuple[str, str]:
    """
    Extract plain text from any uploaded file.
    Returns (extracted_text, error_message).
    """
    import io

    name = uploaded_file.name
    ext  = os.path.splitext(name)[1].lower()

    content_bytes = uploaded_file.read()

    if ext in ('.txt', '.md', '.rst', '.log', '.tsv', '.csv'):
        for enc in ('utf-8', 'utf-8-sig', 'latin-1', 'cp1252'):
            try:
                return content_bytes.decode(enc), ''
            except UnicodeDecodeError:
                continue
        return content_bytes.decode('utf-8', errors='replace'), ''

    if ext in ('.json', '.jsonl'):
        try:
            return content_bytes.decode('utf-8', errors='replace'), ''
        except Exception as e:
            return '', str(e)

    if ext in ('.html', '.htm', '.xml'):
        try:
            import re as _re
            raw  = content_bytes.decode('utf-8', errors='replace')
            text = _re.sub(r'<[^>]+>', ' ', raw)
            text = _re.sub(r'\s+', ' ', text).strip()
            return text, ''
        except Exception as e:
            return '', str(e)

    if ext == '.pdf':
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content_bytes))
            pages  = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    pages.append(t)
            return '\n\n'.join(pages), ''
        except ImportError:
            pass
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(content_bytes)) as pdf:
                pages = [p.extract_text() or '' for p in pdf.pages]
            return '\n\n'.join(pages), ''
        except Exception as e:
            return '', f'PDF extraction failed: {e}'

    if ext in ('.docx', '.doc'):
        try:
            import docx as python_docx
            doc = python_docx.Document(io.BytesIO(content_bytes))
            return '\n'.join(p.text for p in doc.paragraphs if p.text.strip()), ''
        except Exception as e:
            return '', f'DOCX extraction failed: {e}'

    if ext in ('.xlsx', '.xlsm', '.xls', '.ods'):
        try:
            import pandas as pd
            engine = 'openpyxl' if ext in ('.xlsx', '.xlsm') else ('xlrd' if ext == '.xls' else 'odf')
            df = pd.read_excel(io.BytesIO(content_bytes), engine=engine)
            return df.to_csv(index=False), ''
        except Exception as e:
            return '', f'Spreadsheet extraction failed: {e}'

    if ext in ('.pptx', '.ppt'):
        try:
            from pptx import Presentation
            prs    = Presentation(io.BytesIO(content_bytes))
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
                slides.append(f"[Slide {i}]\n" + '\n'.join(t for t in texts if t.strip()))
            return '\n\n'.join(slides), ''
        except Exception as e:
            return '', f'PPTX extraction failed: {e}'

    if ext == '.rtf':
        try:
            from striprtf.striprtf import rtf_to_text
            raw = content_bytes.decode('utf-8', errors='replace')
            return rtf_to_text(raw), ''
        except ImportError:
            pass
        import re as _re
        raw  = content_bytes.decode('utf-8', errors='replace')
        text = _re.sub(r'\\[a-z]+\d*\s?', ' ', raw)
        text = _re.sub(r'[{}\\]', '', text)
        return text.strip(), ''

    try:
        return content_bytes.decode('utf-8', errors='replace'), ''
    except Exception as e:
        return '', f'Could not read file: {e}'


def _generate_summary(text: str) -> str:
    """Delegate to the fine-tuned T5 model in core/summarizer.py."""
    from .summarizer import summarize as t5_summarize
    return t5_summarize(text)


# ─── History Views ────────────────────────────────────────────────────────────

@login_required
def history_view(request):
    user_id = request.session['user_id']
    page    = int(request.GET.get('page', 1))
    limit   = settings.ITEMS_PER_PAGE
    skip    = (page - 1) * limit
    history = get_user_history(user_id, skip=skip, limit=limit)
    total   = count_user_history(user_id)
    total_pages = (total + limit - 1) // limit

    for item in history:
        item['id'] = str(item['_id'])

    return render(request, 'core/history.html', {
        'history':     history,
        'total':       total,
        'page':        page,
        'total_pages': total_pages,
    })


# ─── Profile Views ────────────────────────────────────────────────────────────

@login_required
@csrf_protect
def profile_view(request):
    user_id = request.session['user_id']
    user    = get_user_by_id(user_id)
    error   = None
    success_msg = None
    form = ProfileForm(initial={
        'full_name': user['profile'].get('full_name', ''),
        'bio':       user['profile'].get('bio', ''),
    })

    if request.method == 'POST':
        form = ProfileForm(request.POST)
        if form.is_valid():
            update_user_profile(user_id, form.cleaned_data)
            request.session['full_name'] = form.cleaned_data['full_name']
            success_msg = "Profile updated successfully!"
            user = get_user_by_id(user_id)
        else:
            error = "Please correct the errors below."

    history = get_user_history(user_id, limit=5)
    for item in history:
        item['id'] = str(item['_id'])
    return render(request, 'core/profile.html', {
        'user':        user,
        'form':        form,
        'error':       error,
        'success_msg': success_msg,
        'history':     history,
    })


# ─── API Endpoints ────────────────────────────────────────────────────────────

@login_required
def api_history(request):
    """Return user history as JSON for dynamic sidebar."""
    user_id = request.session['user_id']
    history = get_user_history(user_id, limit=20)
    data = []
    for item in history:
        data.append({
            'id':         str(item['_id']),
            'filename':   item['filename'],
            'action':     item['action'],
            'file_size':  item['file_size'],
            'created_at': item['created_at'].isoformat(),
            'result':     item.get('result', '')[:100],
        })
    return JsonResponse({'history': data})


# ─── Index redirect ───────────────────────────────────────────────────────────

def index(request):
    if request.session.get('user_id'):
        return redirect('/dashboard/')
    return redirect('/login/')


# ─── User: Delete Single History Item ────────────────────────────────────────

@login_required
@require_POST
def delete_history_item_view(request, item_id):
    user_id = request.session['user_id']
    user    = get_user_by_id(user_id)

    if delete_history_item(user_id, item_id):
        save_audit_log(
            user_id  = user_id,
            username = user['username'],
            email    = user['email'],
            action   = 'delete_history_item',
            detail   = f"Deleted history item ID: {item_id}"
        )
        messages.success(request, "History item deleted.")
    else:
        messages.error(request, "Could not delete item.")

    next_url = request.POST.get('next', '/history/')
    return redirect(next_url)


# ─── User: Clear All History ──────────────────────────────────────────────────

@login_required
@require_POST
def clear_history_view(request):
    user_id = request.session['user_id']
    user    = get_user_by_id(user_id)
    count   = clear_user_history(user_id)

    save_audit_log(
        user_id  = user_id,
        username = user['username'],
        email    = user['email'],
        action   = 'clear_all_history',
        detail   = f"Cleared {count} history items."
    )
    messages.success(request, f"All {count} history items cleared.")
    return redirect('/history/')


# ─── User: Delete Own Account ─────────────────────────────────────────────────

@login_required
@csrf_protect
def delete_account_view(request):
    user_id = request.session['user_id']
    user    = get_user_by_id(user_id)

    if request.method == 'POST':
        confirm  = request.POST.get('confirm_delete', '').strip()
        password = request.POST.get('password', '').strip()

        if confirm.lower() != user['username'].lower():
            messages.error(request, "Confirmation text does not match your username.")
            return redirect('/profile/')

        from .db import verify_password
        if not verify_password(password, user['password_hash'], user['salt']):
            messages.error(request, "Incorrect password. Account not deleted.")
            return redirect('/profile/')

        if delete_own_account(user_id):
            request.session.flush()
            return redirect('/account-deleted/')
        else:
            messages.error(request, "Could not delete account. Please try again.")
            return redirect('/profile/')

    return redirect('/profile/')


def account_deleted_view(request):
    """Confirmation page shown after account deletion."""
    return render(request, 'core/account_deleted.html')


# ─── Admin: Audit Logs ────────────────────────────────────────────────────────

@admin_required
def admin_audit_logs(request):
    page  = int(request.GET.get('page', 1))
    limit = 20
    skip  = (page - 1) * limit
    logs  = get_all_audit_logs(skip=skip, limit=limit)
    for log in logs:
        log['id'] = str(log['_id'])
    total       = count_audit_logs()
    total_pages = (total + limit - 1) // limit

    return render(request, 'core/admin_audit_logs.html', {
        'logs':        logs,
        'total':       total,
        'page':        page,
        'total_pages': total_pages,
    })


# ─── Audio Summarize View ─────────────────────────────────────────────────────

@login_required
@csrf_protect
def audio_summarize_view(request):
    user_id = request.session['user_id']
    result  = None
    error   = None

    if request.method == 'POST' and request.FILES.get('audio_file'):
        uploaded   = request.FILES['audio_file']
        ext        = os.path.splitext(uploaded.name)[1].lower()
        audio_exts = {'.mp3', '.wav', '.m4a', '.flac', '.ogg', '.aac'}

        if ext in audio_exts:
            # ── Audio → Whisper → T5 ─────────────────────────────────────────
            import tempfile
            from .audio_processor import process_audio_file
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                for chunk in uploaded.chunks(chunk_size=1024 * 1024):
                    tmp.write(chunk)
                tmp_path = tmp.name
            try:
                result = process_audio_file(tmp_path, uploaded.name)
            finally:
                try:
                    os.unlink(tmp_path)
                except Exception:
                    pass
            if result.get('success'):
                save_upload_history(
                    user_id   = user_id,
                    filename  = uploaded.name,
                    file_hash = result.get('file_hash', 'unknown'),
                    file_size = uploaded.size,
                    action    = 'audio_summarize',
                    result    = result['summary'][:200] + ('...' if len(result['summary']) > 200 else ''),
                )
            else:
                error = result.get('error', 'Audio processing failed.')

        else:
            # ── Text/document → extract text → audio T5 ──────────────────────
            from .audio_processor import summarize_text_with_model
            extracted_text, extract_error = _extract_text_from_file(uploaded)
            if extract_error:
                error = f'Could not read "{uploaded.name}": {extract_error}'
            elif not extracted_text.strip():
                error = f'No text could be extracted from "{uploaded.name}".'
            else:
                summary   = summarize_text_with_model(extracted_text)
                file_hash = compute_sha256(extracted_text.encode())
                result    = {
                    'success':     True,
                    'transcript':  extracted_text[:1000],
                    'summary':     summary,
                    'word_count':  len(extracted_text.split()),
                    'compression': f"{100 - (len(summary.split()) / max(len(extracted_text.split()), 1) * 100):.0f}%",
                }
                save_upload_history(
                    user_id   = user_id,
                    filename  = uploaded.name,
                    file_hash = file_hash,
                    file_size = uploaded.size,
                    action    = 'audio_summarize',
                    result    = summary[:200] + ('...' if len(summary) > 200 else ''),
                )

    history = get_user_history(user_id, limit=5)
    for item in history:
        item['id'] = str(item['_id'])
    return render(request, 'core/audio_summarize.html', {
        'result':  result,
        'error':   error,
        'history': history,
    })


# ─── Video Summarize View ─────────────────────────────────────────────────────

@login_required
@csrf_protect
def video_summarize_view(request):
    user_id = request.session['user_id']
    result  = None
    error   = None

    if request.method == 'POST' and request.FILES.get('video_file'):
        uploaded   = request.FILES['video_file']
        ext        = os.path.splitext(uploaded.name)[1].lower()
        video_exts = {'.mp4', '.avi', '.mov', '.mkv', '.webm', '.flv', '.wmv',
                      '.mp3', '.wav', '.m4a', '.flac', '.ogg'}

        if ext in video_exts:
            # ── Video/Audio → FFmpeg → Whisper → T5 ──────────────────────────
            import tempfile
            from .video_processor import process_video_file
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                for chunk in uploaded.chunks(chunk_size=1024 * 1024):
                    tmp.write(chunk)
                tmp_path = tmp.name
            try:
                result = process_video_file(tmp_path, uploaded.name)
            finally:
                try:
                    os.unlink(tmp_path)
                except Exception:
                    pass
            if result.get('success'):
                save_upload_history(
                    user_id   = user_id,
                    filename  = uploaded.name,
                    file_hash = result.get('file_hash', 'unknown'),
                    file_size = uploaded.size,
                    action    = 'video_summarize',
                    result    = result['summary'][:200] + ('...' if len(result['summary']) > 200 else ''),
                )
            else:
                error = result.get('error', 'Video processing failed.')

        else:
            # ── Text/document → extract text → video T5 ──────────────────────
            from .video_processor import summarize_text_with_model
            extracted_text, extract_error = _extract_text_from_file(uploaded)
            if extract_error:
                error = f'Could not read "{uploaded.name}": {extract_error}'
            elif not extracted_text.strip():
                error = f'No text could be extracted from "{uploaded.name}".'
            else:
                summary   = summarize_text_with_model(extracted_text)
                file_hash = compute_sha256(extracted_text.encode())
                result    = {
                    'success':     True,
                    'transcript':  extracted_text[:1000],
                    'summary':     summary,
                    'word_count':  len(extracted_text.split()),
                    'compression': f"{100 - (len(summary.split()) / max(len(extracted_text.split()), 1) * 100):.0f}%",
                }
                save_upload_history(
                    user_id   = user_id,
                    filename  = uploaded.name,
                    file_hash = file_hash,
                    file_size = uploaded.size,
                    action    = 'video_summarize',
                    result    = summary[:200] + ('...' if len(summary) > 200 else ''),
                )

    history = get_user_history(user_id, limit=5)
    for item in history:
        item['id'] = str(item['_id'])
    return render(request, 'core/video_summarize.html', {
        'result':  result,
        'error':   error,
        'history': history,
    })